"""office_worker: JSON-stdio bridge for local office document writes.

Local fork (2026-09-02) — runs inside the dsh-doc offline Python runtime.
Reads one JSON request from stdin, writes one JSON result to stdout.
Ops:
  xlsx_write {path, mode: create|update, sheets: [{name, rows: [[...]]}]}
  docx_write {path, mode: create|append|replace, blocks?: [...], pairs?: [[old,new]]}
"""
import json
import sys


def main():
    try:
        raw = sys.stdin.buffer.read()
        req = json.loads(raw.decode('utf-8'))
        op = req.get('op')
        if op == 'xlsx_write':
            result = xlsx_write(req)
        elif op == 'docx_write':
            result = docx_write(req)
        else:
            result = {'ok': False, 'error': 'unknown op: %s' % op}
    except Exception as exc:  # noqa: BLE001 — worker boundary must always answer
        result = {'ok': False, 'error': '%s: %s' % (type(exc).__name__, str(exc))}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode('utf-8'))
    sys.stdout.buffer.flush()


def xlsx_write(req):
    import openpyxl
    path = req['path']
    mode = req.get('mode', 'create')
    sheets = req.get('sheets') or []
    if mode == 'create':
        wb = openpyxl.Workbook()
        if sheets:
            wb.remove(wb.active)
    else:
        wb = openpyxl.load_workbook(path)
    for spec in sheets:
        name = str(spec.get('name') or 'Sheet1')[:31]
        if name in wb.sheetnames:
            ws = wb[name]
        else:
            ws = wb.create_sheet(name)
        rows = spec.get('rows') or []
        if not rows:
            continue
        for r, row in enumerate(rows):
            if not isinstance(row, list):
                row = [row]
            for c, value in enumerate(row):
                if value is None:
                    continue
                # Write values verbatim (strings keep leading zeros; existing
                # cell styles survive value assignment in openpyxl).
                ws.cell(row=r + 1, column=c + 1).value = value
    wb.save(path)
    return {'ok': True, 'sheets': wb.sheetnames}


def docx_write(req):
    from docx import Document
    path = req['path']
    mode = req.get('mode', 'create')
    doc = Document() if mode == 'create' else Document(path)
    if mode == 'replace':
        pairs = [(str(old), str(new)) for old, new in (req.get('pairs') or [])]
        replace_all(doc, pairs)
    else:
        for block in (req.get('blocks') or []):
            kind = block.get('type')
            if kind == 'heading':
                level = int(block.get('level') or 1)
                doc.add_heading(str(block.get('text') or ''), level=max(0, min(9, level)))
            elif kind == 'table':
                rows = block.get('rows') or []
                if rows:
                    ncols = max((len(r) if isinstance(r, list) else 1) for r in rows)
                    table = doc.add_table(rows=0, cols=ncols)
                    table.style = 'Table Grid'
                    for row in rows:
                        cells = table.add_row().cells
                        items = row if isinstance(row, list) else [row]
                        for i, value in enumerate(items):
                            cells[i].text = str(value)
            elif kind == 'bullet':
                doc.add_paragraph(str(block.get('text') or ''), style='List Bullet')
            else:
                doc.add_paragraph(str(block.get('text') or ''))
    doc.save(path)
    return {'ok': True, 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables)}


def replace_all(doc, pairs):
    def replace_paragraphs(paragraphs):
        for para in paragraphs:
            text = para.text
            if not text:
                continue
            new_text = text
            for old, new in pairs:
                new_text = new_text.replace(old, new)
            if new_text == text:
                continue
            for run in para.runs:
                run.text = ''
            if para.runs:
                para.runs[0].text = new_text
            else:
                para.add_run(new_text)

    replace_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_paragraphs(cell.paragraphs)


if __name__ == '__main__':
    main()
